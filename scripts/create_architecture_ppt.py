#!/usr/bin/env python3
"""Create an editable PowerPoint infographic for the GraphWeather architectures."""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GraphWeather_architecture_infographics.pptx"

SW, SH = 13.333, 7.5
FONT = "Arial"
TITLE_FONT = "Cambria"
MONO = "Courier New"

C = {
    "bg": "FFFFFF",
    "panel": "FFFFFF",
    "panel2": "F4F4F4",
    "panel3": "E8EEF7",
    "white": "111111",
    "on_dark": "FFFFFF",
    "muted": "555555",
    "line": "777777",
    "black": "25282D",
    "charcoal": "292D33",
    "down_fill": "DDE6F3",
    "up_fill": "F4EEE5",
    "cyan": "2F80ED",
    "blue": "5C7CAC",
    "orange": "F07C42",
    "green": "319B63",
    "purple": "9B7653",
    "red": "D6525D",
    "yellow": "1769E0",
    "h128": "2F80ED",
    "h160": "F07C42",
}


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_bg(slide, color: str = C["bg"]):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def rect(slide, x, y, w, h, fill=C["panel"], line=C["line"], radius=True, lw=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(lw)
    return shape


def textbox(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=12,
    color=C["white"],
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    font=FONT,
    margin=0.03,
    rotation=0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.rotation = rotation
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def rich_textbox(slide, runs, x, y, w, h, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.03):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    for spec in runs:
        run = p.add_run()
        run.text = spec[0]
        run.font.name = spec[4] if len(spec) > 4 else FONT
        run.font.size = Pt(spec[1])
        run.font.bold = spec[2]
        run.font.color.rgb = rgb(spec[3])
    return shape


def line(slide, x1, y1, x2, y2, color=C["line"], width=1.4, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def circle(slide, cx, cy, diameter, fill=C["cyan"], line_color=None, lw=0.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(cx - diameter / 2),
        Inches(cy - diameter / 2),
        Inches(diameter),
        Inches(diameter),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color or fill)
    shape.line.width = Pt(lw)
    return shape


def title(slide, heading, subtitle=None, accent=C["cyan"]):
    textbox(slide, heading, 0.50, 0.30, 12.30, 0.60, 28, C["white"], True, font=TITLE_FONT)
    if subtitle:
        textbox(slide, subtitle, 0.50, 0.90, 12.30, 0.40, 13, C["muted"])


def footer(slide, source, page):
    textbox(slide, source, 0.50, 7.23, 11.8, 0.16, 6.0, "888888", font=MONO)
    textbox(slide, f"{page:02d}", 12.35, 7.21, 0.48, 0.18, 6.5, "888888", True, PP_ALIGN.RIGHT)


def pill(slide, text, x, y, w, color=C["cyan"], fill=C["panel2"], size=8.5):
    rect(slide, x, y, w, 0.28, fill, color, True, 0.8)
    textbox(slide, text, x + 0.03, y + 0.01, w - 0.06, 0.24, size, color, True, PP_ALIGN.CENTER)


def stage_card(slide, x, y, w, h, heading, body, accent, fill=C["panel2"], body_size=8.0):
    if accent == C["cyan"]:
        fill, edge, heading_color, body_color = C["down_fill"], C["black"], C["black"], C["black"]
    elif accent == C["purple"]:
        fill, edge, heading_color, body_color = C["up_fill"], C["black"], C["black"], C["black"]
    elif accent == C["yellow"]:
        fill, edge, heading_color, body_color = C["yellow"], C["black"], C["on_dark"], C["on_dark"]
    else:
        edge, heading_color, body_color = C["black"], C["black"], C["black"]
    rect(slide, x, y, w, h, fill, edge, True, 1.2)
    textbox(slide, heading, x + 0.10, y + 0.06, w - 0.20, 0.23, 10.2, heading_color, True, PP_ALIGN.CENTER)
    textbox(slide, body, x + 0.10, y + 0.29, w - 0.20, h - 0.34, body_size, body_color, False, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)


def process_card(slide, x, y, w, h, heading, body, accent, size=8.2):
    rect(slide, x, y, w, h, C["panel"], accent, True, 1.2)
    textbox(slide, heading, x + 0.06, y + 0.06, w - 0.12, 0.22, 9.2, accent, True, PP_ALIGN.CENTER)
    textbox(slide, body, x + 0.07, y + 0.29, w - 0.14, h - 0.34, size, C["white"], False, PP_ALIGN.CENTER)


def outer_card(slide, x, y, w, h, heading, body, size=7.8):
    rect(slide, x, y, w, h, C["charcoal"], C["black"], True, 1.2)
    textbox(slide, heading, x + 0.06, y + 0.06, w - 0.12, 0.22, 9.2, C["on_dark"], True, PP_ALIGN.CENTER)
    textbox(slide, body, x + 0.07, y + 0.29, w - 0.14, h - 0.34, size, C["on_dark"], False, PP_ALIGN.CENTER)


def connect_u(slide, cards_left, bottom, cards_right, pool_text, unpool_text):
    # Behind-card directional spine.
    for a, b, label in zip(cards_left[:-1], cards_left[1:], pool_text[:-1]):
        x1, y1, w1, h1 = a
        x2, y2, w2, _ = b
        line(slide, x1 + w1 / 2, y1 + h1, x2 + w2 / 2, y2, C["black"], 1.5)
        textbox(slide, f"↓ {label}", (x1 + x2) / 2 - 0.22, (y1 + h1 + y2) / 2 - 0.11, 0.75, 0.2, 7.0, C["muted"], False, PP_ALIGN.CENTER)
    a = cards_left[-1]
    x1, y1, w1, h1 = a
    xb, yb, wb, _ = bottom
    line(slide, x1 + w1 / 2, y1 + h1, xb + wb / 2, yb, C["black"], 1.5)
    textbox(slide, f"↓ {pool_text[-1]}", (x1 + xb) / 2 - 0.22, (y1 + h1 + yb) / 2 - 0.11, 0.75, 0.2, 7.0, C["muted"], False, PP_ALIGN.CENTER)

    a = bottom
    b = cards_right[-1]
    x1, y1, w1, h1 = a
    x2, y2, w2, h2 = b
    line(slide, x1 + w1 / 2, y1, x2 + w2 / 2, y2 + h2, C["black"], 1.5)
    textbox(slide, f"↑ {unpool_text[-1]}", (x1 + x2) / 2 - 0.23, (y1 + y2 + h2) / 2 - 0.11, 0.85, 0.2, 7.0, C["muted"], False, PP_ALIGN.CENTER)
    rev = list(reversed(cards_right))
    for lower, upper, label in zip(rev[:-1], rev[1:], reversed(unpool_text[:-1])):
        xl, yl, wl, _ = lower
        xu, yu, wu, hu = upper
        line(slide, xl + wl / 2, yl, xu + wu / 2, yu + hu, C["black"], 1.5)
        textbox(slide, f"↑ {label}", (xl + xu) / 2 - 0.23, (yl + yu + hu) / 2 - 0.11, 0.85, 0.2, 7.0, C["muted"], False, PP_ALIGN.CENTER)


def add_skips(slide, left_cards, right_cards):
    for left_card, right_card in zip(left_cards, right_cards):
        xl, yl, wl, hl = left_card
        xr, yr, _, hr = right_card
        y = yl + hl / 2
        line(slide, xl + wl, y, xr, yr + hr / 2, C["orange"], 1.1, MSO_LINE_DASH_STYLE.DASH)
        textbox(slide, "skip fusion", (xl + wl + xr) / 2 - 0.35, y - 0.13, 0.70, 0.20, 6.8, C["orange"], False, PP_ALIGN.CENTER)


def add_icosphere(slide, x, y, w, h, accent=C["green"]):
    phi = (1.0 + math.sqrt(5.0)) / 2.0
    verts = []
    for a in (-1, 1):
        for b in (-phi, phi):
            verts.append((0, a, b))
            verts.append((a, b, 0))
            verts.append((b, 0, a))
    verts = np.asarray(verts, dtype=float)
    verts /= np.linalg.norm(verts, axis=1, keepdims=True)
    # Deduplicate and construct native edges by nearest pair distance.
    verts = np.unique(np.round(verts, 9), axis=0)
    ay, az = math.radians(-18), math.radians(14)
    ry = np.array([[math.cos(ay), 0, math.sin(ay)], [0, 1, 0], [-math.sin(ay), 0, math.cos(ay)]])
    rz = np.array([[math.cos(az), -math.sin(az), 0], [math.sin(az), math.cos(az), 0], [0, 0, 1]])
    v = verts @ (rz @ ry).T
    d = np.linalg.norm(v[:, None, :] - v[None, :, :], axis=-1)
    edge_len = np.min(d[d > 1e-6])
    edges = np.argwhere((d > 1e-6) & (d < edge_len * 1.05))
    edges = [(int(i), int(j)) for i, j in edges if i < j]
    pxy = np.c_[x + w * (0.5 + 0.43 * v[:, 0]), y + h * (0.5 - 0.43 * v[:, 1])]
    circle(slide, x + w / 2, y + h / 2, min(w, h) * 0.92, C["panel"], C["line"], 1.0)
    for i, j in sorted(edges, key=lambda e: (v[e[0], 2] + v[e[1], 2])):
        front = (v[i, 2] + v[j, 2]) > 0
        line(slide, pxy[i, 0], pxy[i, 1], pxy[j, 0], pxy[j, 1], accent if front else C["line"], 1.2 if front else 0.55)
    for i, (px, py) in enumerate(pxy):
        if v[i, 2] > -0.1:
            circle(slide, px, py, 0.045, accent, accent, 0)


def mini_u(slide, x, y, w, h, levels, accent, mesh=False):
    xs_l = [x + w * (0.08 + i * 0.115) for i in range(levels - 1)]
    xs_r = [x + w * (0.80 - i * 0.115) for i in range(levels - 1)]
    ys = [y + h * (0.08 + i * 0.18) for i in range(levels - 1)]
    xb, yb = x + w * 0.45, y + h * (0.08 + (levels - 1) * 0.18)
    for i in range(levels - 2):
        line(slide, xs_l[i] + 0.14, ys[i] + 0.14, xs_l[i + 1] + 0.14, ys[i + 1] + 0.14, accent, 1.3)
        line(slide, xs_r[i] + 0.14, ys[i] + 0.14, xs_r[i + 1] + 0.14, ys[i + 1] + 0.14, accent, 1.3)
    line(slide, xs_l[-1] + 0.14, ys[-1] + 0.14, xb + 0.14, yb + 0.14, accent, 1.3)
    line(slide, xb + 0.14, yb + 0.14, xs_r[-1] + 0.14, ys[-1] + 0.14, accent, 1.3)
    for px, py in list(zip(xs_l, ys)) + list(zip(xs_r, ys)) + [(xb, yb)]:
        if mesh:
            circle(slide, px + 0.14, py + 0.14, 0.25, C["panel2"], accent, 1.1)
        else:
            rect(slide, px, py, 0.28, 0.28, C["panel2"], accent, True, 1.0)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    textbox(slide, "Model Architectures", 0.75, 1.50, 11.80, 1.00, 40, C["white"], True, font=TITLE_FONT)
    textbox(slide, "A graph U-Net (encode–process–decode) at three settings — 2.5°, 1.5°, and the GraphCast-style icosphere mesh", 0.75, 2.55, 11.80, 0.50, 16, "444444")
    cards = [
        ("2.5° · L3 (4 levels)", "hidden 128 → 2.47M  ·  hidden 160 → 3.84M", "FFFFFF"),
        ("1.5° · L4 (5 levels)", "hidden 160 → 4.60M  ·  EMA weights", "FFFFFF"),
        ("Icosphere · M5→M2", "hidden 160 → 3.84M  ·  native k=6", "F4F4F4"),
    ]
    for i, (h, b, fill) in enumerate(cards):
        x = 0.68 + i * 4.10
        rect(slide, x, 4.25, 3.75, 1.50, fill, C["black"], True, 1.1)
        textbox(slide, h, x + 0.29, 4.53, 3.15, 0.55, 16, C["white"], True, font=TITLE_FONT)
        textbox(slide, b, x + 0.29, 5.07, 3.15, 0.50, 11.5, "444444")
    footer(slide, "GraphWeather5p625 · architecture configs and model summaries · generated 2026-07-21", 1)


def slide_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Three processor topologies, four trained widths", "The two 2.5° grid experiments share one topology; width alone changes.")
    cols = [
        ("1.5° GRID · H160", "lat-lon", "5", "29,040→120", "8 / 24", "4,596,522", C["blue"], 5, False),
        ("2.5° GRID · H128", "lat-lon", "4", "10,368→162", "8 / 24", "2,473,847", C["h128"], 4, False),
        ("2.5° GRID · H160", "lat-lon", "4", "10,368→162", "8 / 24", "3,844,932", C["h160"], 4, False),
        ("2.5° M5 MESH · H160", "icosphere", "4", "10,242→162", "native 5/6", "3,842,782", C["green"], 4, True),
    ]
    x0, gap, cw = 0.42, 0.16, 3.0
    for i, (name, domain, levels, nodes, degree, params, accent, nlev, mesh) in enumerate(cols):
        x = x0 + i * (cw + gap)
        rect(slide, x, 1.14, cw, 5.55, C["panel"], accent, True, 1.4)
        textbox(slide, name, x + 0.16, 1.35, cw - 0.32, 0.36, 12, accent, True, PP_ALIGN.CENTER)
        mini_u(slide, x + 0.27, 1.83, cw - 0.54, 1.8, nlev, accent, mesh)
        if mesh:
            add_icosphere(slide, x + 1.93, 1.96, 0.72, 0.72, accent)
        rows = [("domain", domain), ("levels", levels), ("nodes", nodes), ("degree k", degree), ("parameters", params)]
        for j, (k, v) in enumerate(rows):
            yy = 3.87 + j * 0.46
            textbox(slide, k.upper(), x + 0.18, yy, 0.9, 0.26, 7.1, C["muted"], True)
            textbox(slide, v, x + 1.10, yy, cw - 1.28, 0.26, 9.5, C["white"], True, PP_ALIGN.RIGHT)
        note = "BipartiteMP only · no outer attention" if mesh else ("extra L4 + return refine" if nlev == 5 else "same graph + same blocks")
        pill(slide, note, x + 0.31, 6.10, cw - 0.62, accent, C["panel2"], 7.4)
    footer(slide, "Sources: the three named experiment configs + runs/2p5_l3_h160_icomeshm5_bipmponly.../config_resolved.yaml", 2)


def slide_1p5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "1.5° lat-lon Graph U-Net · H160", "Five levels; dense k=24 only at the 120-node bottleneck; symmetric skip-fused decoder.", C["blue"])
    pill(slide, "4,596,522 trainable parameters", 10.38, 0.37, 2.45, C["blue"])
    rect(slide, 1.25, 1.18, 10.78, 5.64, C["bg"], C["line"], True, 1.0)

    left = [(1.45, 1.92, 1.85, 0.93), (2.35, 2.94, 1.85, 0.93), (3.25, 3.96, 1.85, 0.93), (4.15, 4.98, 1.85, 0.93)]
    right = [(10.03, 1.92, 1.85, 0.93), (9.13, 2.94, 1.85, 0.93), (8.23, 3.96, 1.85, 0.93), (7.33, 4.98, 1.85, 0.93)]
    bottom = (5.69, 5.91, 1.95, 0.82)
    add_skips(slide, left, right)
    connect_u(slide, left, bottom, right, ["Pool 51.4k"] * 4, ["Unpool 77.1k"] * 4)
    process_card(slide, 0.28, 1.24, 1.00, 1.28, "INPUT", "[B,134]\n121×240", C["cyan"], 8.2)
    outer_card(slide, 1.42, 1.19, 1.90, 0.62, "EMBED + ENCODER", "134→160 · 21.6k  |  Attn×1 · 311.6k", 7.2)
    outer_card(slide, 10.00, 1.19, 1.90, 0.62, "DECODER + HEAD", "Attn×1 · 311.6k  |  160→67 · 10.8k", 7.2)
    pill(slide, "each ×N stage = LocalGraphAttentionBlock · LN→MHA→MLP(4D)", 4.32, 1.34, 4.68, C["yellow"], C["panel2"], 7.1)
    process_card(slide, 12.05, 1.24, 1.00, 1.28, "OUTPUT", "+ current\n[B,67,121,240]", C["orange"], 7.5)
    line(slide, 1.28, 1.88, 1.45, 1.88, C["cyan"], 1.7)
    line(slide, 11.88, 1.88, 12.05, 1.88, C["purple"], 1.7)

    data = [
        ("L0 · DOWN ×2", "121×240 · N=29,040\nk=8 · E=232,320\nD160 · 311.6k/block"),
        ("L1 · DOWN ×2", "61×120 · N=7,320\nk=8 · E=58,560\nD160 · 311.6k/block"),
        ("L2 · DOWN ×1", "31×60 · N=1,860\nk=8 · E=14,880\nD160 · 311.6k/block"),
        ("L3 · DOWN ×1", "16×30 · N=480\nk=8 · E=3,840\nD160 · 311.6k/block"),
    ]
    up = [
        ("L0 · REFINE ×1", data[0][1]),
        ("L1 · REFINE ×1", data[1][1]),
        ("L2 · REFINE ×1", data[2][1]),
        ("L3 · REFINE ×1", data[3][1]),
    ]
    for c, d in zip(left, data): stage_card(slide, *c, d[0], d[1], C["cyan"])
    for c, d in zip(right, up): stage_card(slide, *c, d[0], d[1], C["purple"])
    stage_card(slide, *bottom, "L4 · BOTTLENECK ×1", "8×15 · N=120  |  k=24 · E=2,880\nD160 · 311.6k/block", C["yellow"], C["panel3"], 7.8)
    textbox(slide, "MeanMaxPool: child mean + max → 2D→D", 1.53, 6.62, 3.25, 0.18, 6.8, C["muted"])
    textbox(slide, "ParentUnpoolFuse: parent gather + skip → 2D→D→D", 8.15, 6.62, 3.7, 0.18, 6.8, C["muted"], align=PP_ALIGN.RIGHT)
    footer(slide, "config_1p5_l4_hidden160_dense_l4k24_curriculum_S2toS10_3ep_initckpt_s1_150ep.yaml", 3)


def slide_2p5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "2.5° lat-lon Graph U-Net · H128 versus H160", "Topology is identical. Every colored number below is a width-dependent parameter change.", C["orange"])
    pill(slide, "H128 · 2,473,847", 9.12, 0.36, 1.68, C["h128"])
    pill(slide, "H160 · 3,844,932", 10.96, 0.36, 1.86, C["h160"])
    rect(slide, 1.18, 1.18, 10.95, 5.63, C["bg"], C["line"], True, 1.0)
    left = [(1.50, 2.00, 2.08, 1.00), (2.56, 3.24, 2.08, 1.00), (3.62, 4.48, 2.08, 1.00)]
    right = [(9.75, 2.00, 2.08, 1.00), (8.69, 3.24, 2.08, 1.00), (7.63, 4.48, 2.08, 1.00)]
    bottom = (5.62, 5.69, 2.08, 1.00)
    add_skips(slide, left, right)
    connect_u(slide, left, bottom, right, ["Pool"] * 3, ["Unpool"] * 3)
    process_card(slide, 0.22, 1.35, 1.06, 1.32, "INPUT", "[B,134]\n72×144", C["cyan"], 8.2)
    outer_card(slide, 1.47, 1.20, 2.10, 0.67, "EMBED + ENCODER×1", "H128 17.3k + 200.1k\nH160 21.6k + 311.6k", 7.4)
    outer_card(slide, 9.76, 1.20, 2.10, 0.67, "DECODER×1 + HEAD", "H128 200.1k + 8.6k\nH160 311.6k + 10.8k", 7.4)
    pill(slide, "each ×N stage = LocalGraphAttentionBlock · LN→MHA→MLP(4D)", 4.34, 1.36, 4.64, C["yellow"], C["panel2"], 7.1)
    process_card(slide, 12.05, 1.35, 1.06, 1.32, "OUTPUT", "+ current\n[B,67,72,144]", C["orange"], 7.8)
    line(slide, 1.28, 2.00, 1.50, 2.00, C["cyan"], 1.7)
    line(slide, 11.83, 2.00, 12.05, 2.00, C["purple"], 1.7)
    levels = [
        ("L0", "72×144 · N=10,368 · k8 · E=82,944", "DOWN ×2", "REFINE ×1"),
        ("L1", "36×72 · N=2,592 · k8 · E=20,736", "DOWN ×2", "REFINE ×1"),
        ("L2", "18×36 · N=648 · k8 · E=5,184", "DOWN ×1", "REFINE ×1"),
    ]
    dual = "H128: D128 / 200.1k    |    H160: D160 / 311.6k"
    for c, (lv, meta, down, _) in zip(left, levels):
        stage_card(slide, *c, f"{lv} · {down}", f"{meta}\n{dual}", C["cyan"], body_size=7.2)
    for c, (lv, meta, _, up) in zip(right, levels):
        stage_card(slide, *c, f"{lv} · {up}", f"{meta}\n{dual}", C["purple"], body_size=7.2)
    stage_card(slide, *bottom, "L3 · BOTTLENECK ×1", "9×18 · N=162 · k24 · E=3,888\nH128: 200.1k  |  H160: 311.6k", C["yellow"], C["panel3"], 7.5)
    pill(slide, "Pool: 32.9k → 51.4k", 0.42, 6.45, 1.80, C["cyan"], C["panel2"], 7.4)
    pill(slide, "Attention block: 200.1k → 311.6k", 2.40, 6.45, 2.55, C["orange"], C["panel2"], 7.4)
    pill(slide, "Unpool: 49.4k → 77.1k", 8.34, 6.45, 2.00, C["purple"], C["panel2"], 7.4)
    pill(slide, "+55.4% total parameters", 10.53, 6.45, 2.34, C["h160"], C["panel2"], 7.4)
    footer(slide, "config_2p5_l3_hidden128...initckpt.yaml  vs  config_2p5_l3_hidden160...initckpt_s1_200ep.yaml", 4)


def slide_icosphere(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "2.5° GraphCast-style grid↔icosphere · H160", "The existing Graph U-Net runs on native M5→M2 triangulations; only domain crossing is new.", C["green"])
    pill(slide, "3,842,782 trainable parameters", 10.33, 0.36, 2.50, C["green"])
    add_icosphere(slide, 5.88, 0.78, 1.55, 1.55, C["green"])
    textbox(slide, "native triangular mesh", 5.60, 1.96, 2.10, 0.22, 7.2, C["green"], True, PP_ALIGN.CENTER)
    rect(slide, 1.72, 1.46, 9.88, 5.35, C["bg"], C["line"], True, 1.0)
    left = [(2.12, 2.18, 2.05, 0.98), (3.10, 3.35, 2.05, 0.98), (4.08, 4.52, 2.05, 0.98)]
    right = [(9.18, 2.18, 2.05, 0.98), (8.20, 3.35, 2.05, 0.98), (7.22, 4.52, 2.05, 0.98)]
    bottom = (5.64, 5.65, 2.05, 0.98)
    add_skips(slide, left, right)
    connect_u(slide, left, bottom, right, ["Pool 51.4k"] * 3, ["Unpool 77.1k"] * 3)
    process_card(slide, 0.20, 1.38, 1.31, 1.26, "GRID INPUT", "72×144\nN=10,368\n134→160 · 21.6k", C["cyan"], 7.8)
    process_card(slide, 0.38, 2.86, 1.38, 1.16, "GRID→MESH", "BipartiteMP\nE=16,520\n310.1k", C["green"], 7.8)
    process_card(slide, 11.61, 2.86, 1.38, 1.16, "MESH→GRID", "BipartiteMP\nE=31,104 (=3/grid)\n310.1k", C["green"], 7.4)
    process_card(slide, 11.82, 1.38, 1.31, 1.26, "GRID OUTPUT", "160→67 · 10.8k\n+ current\n[B,67,72,144]", C["orange"], 7.4)
    line(slide, 1.51, 2.01, 2.12, 2.67, C["green"], 1.8)
    line(slide, 1.76, 3.44, 2.12, 2.80, C["green"], 1.8)
    line(slide, 11.23, 2.80, 11.61, 3.44, C["green"], 1.8)
    line(slide, 11.99, 2.86, 12.20, 2.64, C["green"], 1.8)
    pill(slide, "mesh static 4→160 · 800 params", 1.84, 1.66, 2.30, C["green"], C["panel2"], 7.2)
    pill(slide, "encoder_blocks = decoder_blocks = 0", 9.04, 1.68, 2.38, C["orange"], C["panel2"], 7.2)
    data = [
        ("M5 / L0 · DOWN ×2", "N=10,242 · k=6 · E=61,452\n20,480 faces · D160\n311.6k/attention block"),
        ("M4 / L1 · DOWN ×2", "N=2,562 · k=6 · E=15,372\n5,120 faces · D160\n311.6k/attention block"),
        ("M3 / L2 · DOWN ×1", "N=642 · k=6 · E=3,852\n1,280 faces · D160\n311.6k/attention block"),
    ]
    up = [
        ("M5 / L0 · REFINE ×1", data[0][1]),
        ("M4 / L1 · REFINE ×1", data[1][1]),
        ("M3 / L2 · REFINE ×1", data[2][1]),
    ]
    for c, d in zip(left, data): stage_card(slide, *c, d[0], d[1], C["cyan"], body_size=7.25)
    for c, d in zip(right, up): stage_card(slide, *c, d[0], d[1], C["purple"], body_size=7.25)
    stage_card(slide, *bottom, "M2 / L3 · BOTTLENECK ×1", "N=162 · k=6 · E=972 · 320 faces\n12 pentagons use one masked dummy slot", C["yellow"], C["panel3"], 7.4)
    textbox(slide, "No kNN mesh edges: connectivity comes from the icosphere triangulation itself.", 3.58, 6.70, 6.2, 0.18, 7.0, C["green"], True, PP_ALIGN.CENTER)
    footer(slide, "runs/2p5_l3_h160_icomeshm5_bipmponly_s1x150_currs2tos10x3/config_resolved.yaml", 5)


def slide_attention(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Block anatomy · LocalGraphAttentionBlock", "The same residual attention+MLP block is reused at every U-Net stage.", C["cyan"])
    # Main residual flow.
    y, h = 1.35, 0.68
    boxes = [
        (0.35, 0.78, "h", "[B,N,D]", C["white"]),
        (1.36, 0.72, "LN", "D", C["cyan"]),
        (2.32, 2.08, "EDGE-AWARE LOCAL ATTENTION", "Q/K/V + edge key/value/bias", C["cyan"]),
        (4.65, 0.70, "+", "residual", C["orange"]),
        (5.58, 0.72, "LN", "D", C["purple"]),
        (6.55, 2.10, "NODE MLP", "D → 4D → D · GELU", C["purple"]),
        (8.90, 0.70, "+", "residual", C["orange"]),
        (9.84, 1.05, "h′", "[B,N,D]", C["white"]),
    ]
    for x, w, hd, bd, a in boxes:
        process_card(slide, x, y, w, h, hd, bd, a, 7.6)
    for (x, w, *_), (x2, *_rest) in zip(boxes[:-1], boxes[1:]):
        line(slide, x + w, y + h / 2, x2, y + h / 2, C["line"], 1.5)
    # residual bypasses
    line(slide, 0.74, y, 0.74, 1.08, C["orange"], 1.1)
    line(slide, 0.74, 1.08, 5.00, 1.08, C["orange"], 1.1)
    line(slide, 5.00, 1.08, 5.00, y, C["orange"], 1.1)
    line(slide, 5.00, y + h, 5.00, 2.25, C["orange"], 1.1)
    line(slide, 5.00, 2.25, 9.25, 2.25, C["orange"], 1.1)
    line(slide, 9.25, 2.25, 9.25, y + h, C["orange"], 1.1)

    # Attention detail.
    rect(slide, 0.45, 2.72, 7.63, 3.83, C["panel"], C["cyan"], True, 1.2)
    textbox(slide, "EDGE-AWARE MULTI-HEAD ATTENTION", 0.70, 2.93, 4.2, 0.28, 12, C["cyan"], True)
    pill(slide, "softmax over k incoming neighbors", 5.45, 2.92, 2.30, C["cyan"], C["panel2"], 7.5)
    attention_rows = [
        ("Q", "Linear D→D", "[B,N,H,32]", C["cyan"]),
        ("K / V", "Linear D→D", "gather [B,N,k,H,32]", C["blue"]),
        ("edge K / V", "Linear 6→D", "[N,k,H,32]", C["green"]),
        ("edge bias", "Linear 6→H", "[N,k,H]", C["orange"]),
    ]
    for i, (name, op, shape, a) in enumerate(attention_rows):
        yy = 3.46 + i * 0.53
        pill(slide, name, 0.72, yy, 1.05, a, C["panel2"], 8.0)
        textbox(slide, op, 1.93, yy, 1.55, 0.28, 9, C["white"], True)
        textbox(slide, shape, 3.50, yy, 2.45, 0.28, 8.5, C["muted"], font=MONO)
    rect(slide, 6.02, 3.47, 1.72, 2.10, C["panel2"], C["line"], True, 0.8)
    textbox(slide, "scores", 6.17, 3.60, 1.40, 0.23, 9.2, C["white"], True, PP_ALIGN.CENTER)
    textbox(slide, "Q·(K+edgeK)\n────────────\n√32\n+ edge bias", 6.17, 3.88, 1.40, 1.12, 10, C["cyan"], True, PP_ALIGN.CENTER)
    textbox(slide, "masked_fill(−∞)\nmesh dummy slots only", 6.15, 5.08, 1.45, 0.38, 7.1, C["green"], True, PP_ALIGN.CENTER)
    textbox(slide, "Weighted Σ(V + edgeV) → output projection D→D", 0.74, 5.83, 6.85, 0.36, 10, C["white"], True, PP_ALIGN.CENTER)

    # MLP and parameter table.
    rect(slide, 8.35, 2.72, 4.53, 3.83, C["panel"], C["purple"], True, 1.2)
    textbox(slide, "BLOCK PARAMETER COUNT", 8.63, 2.93, 3.95, 0.28, 12, C["purple"], True)
    rows = [
        ("Configuration", "H128", "H160"),
        ("heads × head_dim", "4×32", "5×32"),
        ("attention + LN₁", "68,124", "105,635"),
        ("MLP + LN₂", "131,968", "205,920"),
        ("full block", "200,092", "311,555"),
    ]
    for i, row in enumerate(rows):
        yy = 3.42 + i * 0.50
        if i == 0:
            rect(slide, 8.60, yy, 3.98, 0.39, C["panel3"], C["line"], False, 0.5)
        textbox(slide, row[0], 8.70, yy, 1.74, 0.34, 8.2, C["muted"] if i else C["white"], i == 0)
        textbox(slide, row[1], 10.43, yy, 0.95, 0.34, 8.8, C["h128"], True, PP_ALIGN.RIGHT)
        textbox(slide, row[2], 11.43, yy, 1.03, 0.34, 8.8, C["h160"], True, PP_ALIGN.RIGHT)
    textbox(slide, "Heads repartition D; Q/K/V/out remain D→D. Fewer heads therefore save almost no parameters.", 8.66, 6.05, 3.80, 0.36, 7.7, C["muted"], False, PP_ALIGN.CENTER)
    footer(slide, "Implementation: src/layers.py · MLP ratio=4 · edge feature dimension=6", 6)


def slide_transfer(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Block anatomy · pooling, unpooling, and symmetric skips", "Resolution changes are learned projections; graph connectivity itself is never rewritten.", C["purple"])
    # Pool panel.
    rect(slide, 0.45, 1.25, 5.98, 4.05, C["panel"], C["cyan"], True, 1.2)
    textbox(slide, "MeanMaxPool · DOWN", 0.72, 1.47, 3.5, 0.32, 15, C["cyan"], True)
    textbox(slide, "parent map: one coarse index per fine node", 0.73, 1.82, 4.8, 0.24, 8.5, C["muted"])
    # child node illustration
    child_pts = [(1.0, 2.55), (1.45, 2.32), (1.52, 2.78), (1.92, 2.52)]
    for px, py in child_pts: circle(slide, px, py, 0.16, C["cyan"])
    for px, py in child_pts: line(slide, px, py, 2.55, 2.55, C["line"], 0.8)
    circle(slide, 2.55, 2.55, 0.23, C["yellow"])
    textbox(slide, "children [B,Nfine,D]", 0.78, 3.00, 1.75, 0.24, 8.2, C["muted"], True, PP_ALIGN.CENTER)
    process_card(slide, 2.96, 2.12, 1.18, 0.88, "REDUCE", "mean  +  max", C["cyan"], 8.3)
    process_card(slide, 4.48, 2.12, 1.42, 0.88, "PROJECT", "concat 2D\nLinear 2D→D", C["green"], 7.8)
    line(slide, 2.68, 2.55, 2.96, 2.55, C["cyan"], 1.5)
    line(slide, 4.14, 2.55, 4.48, 2.55, C["cyan"], 1.5)
    textbox(slide, "H128: 32,896 params", 1.02, 3.62, 2.15, 0.32, 10, C["h128"], True, PP_ALIGN.CENTER)
    textbox(slide, "H160: 51,360 params", 3.55, 3.62, 2.15, 0.32, 10, C["h160"], True, PP_ALIGN.CENTER)
    pill(slide, "grid: proportional parent index", 0.92, 4.32, 2.30, C["blue"])
    pill(slide, "mesh: refinement parent map", 3.50, 4.32, 2.30, C["green"])

    # Unpool panel.
    rect(slide, 6.72, 1.25, 6.16, 4.05, C["panel"], C["purple"], True, 1.2)
    textbox(slide, "ParentUnpoolFuse · UP", 7.00, 1.47, 3.8, 0.32, 15, C["purple"], True)
    textbox(slide, "gather the parent embedding, then fuse the same-level skip", 7.01, 1.82, 5.1, 0.24, 8.5, C["muted"])
    process_card(slide, 7.18, 2.20, 1.30, 0.90, "PARENT", "gather hcoarse\n[B,Nfine,D]", C["purple"], 7.5)
    process_card(slide, 7.18, 3.23, 1.30, 0.90, "SKIP", "encoder feature\n[B,Nfine,D]", C["cyan"], 7.5)
    process_card(slide, 9.06, 2.57, 1.54, 1.18, "FUSE", "concat 2D\nLinear 2D→D\nGELU · Linear D→D", C["orange"], 7.3)
    line(slide, 8.48, 2.65, 9.06, 2.95, C["purple"], 1.5)
    line(slide, 8.48, 3.68, 9.06, 3.36, C["cyan"], 1.5)
    process_card(slide, 11.16, 2.69, 1.20, 0.92, "REFINE", "Attention×1\non fine graph", C["green"], 7.6)
    line(slide, 10.60, 3.15, 11.16, 3.15, C["purple"], 1.5)
    textbox(slide, "H128: 49,408 params", 7.25, 4.44, 2.15, 0.32, 10, C["h128"], True, PP_ALIGN.CENTER)
    textbox(slide, "H160: 77,120 params", 10.04, 4.44, 2.15, 0.32, 10, C["h160"], True, PP_ALIGN.CENTER)

    # Symmetry strip.
    rect(slide, 0.45, 5.61, 12.43, 1.16, C["panel2"], C["line"], True, 0.8)
    textbox(slide, "DOWN PATH", 0.72, 5.79, 1.15, 0.22, 8.2, C["cyan"], True)
    textbox(slide, "Attention blocks  →  save skip  →  MeanMaxPool", 1.83, 5.72, 3.40, 0.36, 10, C["white"], True, PP_ALIGN.CENTER)
    textbox(slide, "BOTTLENECK", 5.64, 5.72, 1.35, 0.36, 10, C["yellow"], True, PP_ALIGN.CENTER)
    textbox(slide, "ParentUnpoolFuse  →  Attention refine", 7.15, 5.72, 3.40, 0.36, 10, C["white"], True, PP_ALIGN.CENTER)
    textbox(slide, "UP PATH", 11.19, 5.79, 1.10, 0.22, 8.2, C["purple"], True, PP_ALIGN.RIGHT)
    line(slide, 5.20, 5.90, 5.58, 5.90, C["cyan"], 1.6)
    line(slide, 7.00, 5.90, 7.13, 5.90, C["purple"], 1.6)
    textbox(slide, "The drawn U is view-symmetric; block counts need not be numerically mirrored (e.g. L0 down×2, up×1).", 1.35, 6.30, 10.6, 0.22, 7.6, C["muted"], False, PP_ALIGN.CENTER)
    footer(slide, "Implementation: src/pooling.py + src/processor.py · include_max=true · mean_type=mean", 7)


def slide_bipartite(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Block anatomy · BipartiteMP", "A variable-degree GraphCast-style transfer block is used only at grid↔mesh boundaries.", C["green"])
    # Source/destination graph illustration.
    rect(slide, 0.45, 1.20, 12.43, 2.15, C["panel"], C["green"], True, 1.2)
    textbox(slide, "SOURCE NODES", 0.72, 1.42, 1.45, 0.23, 8.5, C["cyan"], True)
    textbox(slide, "DESTINATION NODES", 10.67, 1.42, 1.78, 0.23, 8.5, C["orange"], True, PP_ALIGN.RIGHT)
    src_pts = [(1.05, 2.05), (1.05, 2.55), (1.65, 2.30), (2.10, 1.92), (2.13, 2.70)]
    dst_pts = [(11.10, 2.03), (11.10, 2.58), (11.73, 2.30), (12.22, 1.95), (12.20, 2.68)]
    for p in src_pts: circle(slide, *p, 0.14, C["cyan"])
    for p in dst_pts: circle(slide, *p, 0.16, C["orange"])
    process_card(slide, 2.72, 1.62, 2.44, 1.28, "MESSAGE MLP", "concat [hsrc, hdst, edge6]\n326 → 320 → 160\n156,000 params", C["green"], 8.5)
    process_card(slide, 5.52, 1.62, 1.62, 1.28, "AGGREGATE", "scatter_sum\nindex_add_ by dst\nvariable degree", C["orange"], 8.1)
    process_card(slide, 7.50, 1.62, 2.48, 1.28, "UPDATE MLP + RESIDUAL", "concat [hdst, Σmessage]\n320 → 320 → 160\nhdst + update · 154,080 params", C["purple"], 8.0)
    for p in src_pts[1:4]: line(slide, p[0] + 0.08, p[1], 2.72, 2.25, C["line"], 0.8)
    line(slide, 5.16, 2.25, 5.52, 2.25, C["green"], 1.5)
    line(slide, 7.14, 2.25, 7.50, 2.25, C["orange"], 1.5)
    for p in dst_pts[1:4]: line(slide, 9.98, 2.25, p[0] - 0.08, p[1], C["line"], 0.8)
    textbox(slide, "one BipartiteMP = 310,080 parameters", 4.73, 3.02, 3.75, 0.22, 9, C["green"], True, PP_ALIGN.CENTER)

    # Two uses.
    rect(slide, 0.45, 3.68, 6.02, 2.82, C["panel"], C["cyan"], True, 1.1)
    textbox(slide, "GRID → MESH ENCODER", 0.73, 3.92, 3.5, 0.32, 14, C["cyan"], True)
    process_card(slide, 0.84, 4.48, 1.43, 1.04, "GRID", "10,368 nodes\nhgrid [B,N,160]", C["cyan"], 7.7)
    process_card(slide, 2.63, 4.48, 1.36, 1.04, "EDGES", "radius 0.6×\nE=16,520", C["green"], 7.7)
    process_card(slide, 4.36, 4.48, 1.55, 1.04, "M5 MESH", "10,242 nodes\nstatic 4→160", C["orange"], 7.7)
    line(slide, 2.27, 5.00, 2.63, 5.00, C["green"], 1.4)
    line(slide, 3.99, 5.00, 4.36, 5.00, C["green"], 1.4)
    textbox(slide, "Mesh destination starts from learned static lat/lon features (800 params).", 0.88, 5.82, 5.00, 0.34, 8.1, C["muted"], False, PP_ALIGN.CENTER)

    rect(slide, 6.78, 3.68, 6.10, 2.82, C["panel"], C["orange"], True, 1.1)
    textbox(slide, "MESH → GRID DECODER", 7.06, 3.92, 3.5, 0.32, 14, C["orange"], True)
    process_card(slide, 7.18, 4.48, 1.53, 1.04, "M5 MESH", "10,242 nodes\nhmesh [B,M,160]", C["orange"], 7.5)
    process_card(slide, 9.08, 4.48, 1.38, 1.04, "EDGES", "3 triangle vertices\nE=31,104", C["green"], 7.5)
    process_card(slide, 10.84, 4.48, 1.55, 1.04, "GRID", "10,368 nodes\nhgrid is residual", C["cyan"], 7.5)
    line(slide, 8.71, 5.00, 9.08, 5.00, C["green"], 1.4)
    line(slide, 10.46, 5.00, 10.84, 5.00, C["green"], 1.4)
    textbox(slide, "The original embedded grid is carried around the U-Net and updated residually.", 7.23, 5.82, 5.08, 0.34, 8.1, C["muted"], False, PP_ALIGN.CENTER)
    footer(slide, "Implementation: src/mesh_layers.py · D=160 · edge_dim=6 · mlp_hidden_ratio=2", 8)


def slide_parameters(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Parameter ledger · what actually changed", "Exact trainable counts from the run model summaries; bars use a common 4.8M scale.", C["orange"])
    models = [
        ("2.5° GRID H128", 2473847, {"I/O": 25923, "outer": 400184, "attention": 1800828, "transfer": 246912, "mesh": 0}, C["h128"]),
        ("2.5° GRID H160", 3844932, {"I/O": 32387, "outer": 623110, "attention": 2803995, "transfer": 385440, "mesh": 0}, C["h160"]),
        ("1.5° GRID H160", 4596522, {"I/O": 32387, "outer": 623110, "attention": 3427105, "transfer": 513920, "mesh": 0}, C["blue"]),
        ("2.5° M5 MESH H160", 3842782, {"I/O": 32387, "outer": 0, "attention": 2803995, "transfer": 385440, "mesh": 620960}, C["green"]),
    ]
    category_colors = {"I/O": C["muted"], "outer": C["red"], "attention": C["blue"], "transfer": C["purple"], "mesh": C["green"]}
    # Legend.
    lx = 0.55
    for key, label in [("I/O", "embed + head"), ("outer", "pre/post attention"), ("attention", "U-Net attention"), ("transfer", "pool + unpool"), ("mesh", "mesh crossing")]:
        rect(slide, lx, 1.12, 0.18, 0.18, category_colors[key], category_colors[key], False, 0)
        textbox(slide, label, lx + 0.23, 1.08, 1.42, 0.26, 7.8, C["muted"], True)
        lx += 2.38
    # Bars.
    maxv = 4_800_000
    bar_x, bar_w = 3.25, 8.60
    for i, (name, total, parts, accent) in enumerate(models):
        y = 1.66 + i * 1.04
        textbox(slide, name, 0.55, y, 2.40, 0.30, 10.0, accent, True)
        textbox(slide, f"{total/1e6:.3f}M", 0.55, y + 0.31, 1.22, 0.25, 9.2, C["white"], True)
        x = bar_x
        for key in ["I/O", "outer", "attention", "transfer", "mesh"]:
            value = parts[key]
            if not value:
                continue
            w = bar_w * value / maxv
            rect(slide, x, y + 0.04, w, 0.53, category_colors[key], category_colors[key], False, 0)
            if w > 0.70:
                textbox(slide, f"{value/1e6:.2f}M", x + 0.02, y + 0.16, w - 0.04, 0.22, 7.2, C["bg"], True, PP_ALIGN.CENTER)
            x += w
        line(slide, bar_x, y + 0.66, bar_x + bar_w, y + 0.66, C["line"], 0.5)
        textbox(slide, f"{100*total/maxv:.1f}% of scale", 11.95, y + 0.12, 0.85, 0.25, 7.0, C["muted"], True, PP_ALIGN.RIGHT)

    # Change callouts.
    callouts = [
        ("H128 → H160", "+1,371,085 params", "same graph + same blocks; D and heads change", C["h160"]),
        ("2.5° → 1.5° H160", "+751,590 params", "adds L4 block, L3 return refine, pool34 + unpool43", C["blue"]),
        ("grid H160 → M5 mesh", "−2,150 params", "remove 2 outer attention blocks; add 2 BipartiteMP + mesh init", C["green"]),
    ]
    for i, (hd, delta, body, accent) in enumerate(callouts):
        x = 0.55 + i * 4.18
        rect(slide, x, 5.94, 3.85, 0.90, C["panel2"], accent, True, 1.0)
        textbox(slide, hd, x + 0.15, 6.05, 1.42, 0.22, 8.8, accent, True)
        textbox(slide, delta, x + 1.55, 6.05, 2.12, 0.22, 9.2, C["white"], True, PP_ALIGN.RIGHT)
        textbox(slide, body, x + 0.15, 6.34, 3.55, 0.32, 7.2, C["muted"], False, PP_ALIGN.CENTER)
    footer(slide, "Exact totals: model_summary.txt for 1p5 H160, 2p5 H128, 2p5 H160, and 2p5 M5 bipartite-only", 9)


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "GraphWeather Architecture Infographics"
    prs.core_properties.subject = "1.5-degree, 2.5-degree, and GraphCast-style icosphere Graph U-Net architectures"
    prs.core_properties.author = "OpenAI Codex with GraphWeather5p625"
    prs.core_properties.keywords = "GraphWeather, Graph U-Net, icosphere, GraphCast, weather forecasting"
    slide_cover(prs)
    slide_overview(prs)
    slide_1p5(prs)
    slide_2p5(prs)
    slide_icosphere(prs)
    slide_attention(prs)
    slide_transfer(prs)
    slide_bipartite(prs)
    slide_parameters(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
