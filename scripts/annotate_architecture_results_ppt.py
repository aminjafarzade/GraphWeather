#!/usr/bin/env python3
"""Append evidence-based annotations to the existing architecture/results deck.

The source deck is treated as immutable: no slide, shape, or text is removed or
edited.  New callouts are appended only in unused areas of existing slides.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# Anchored at the repo root so the script works from any CWD and survives a
# checkout move (the decks now live in archive/).
REPO_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_INPUT = REPO_ROOT / "archive" / "model_architectures_with_dense_l3k24 (1).pptx"
DEFAULT_OUTPUT = (
    REPO_ROOT / "archive" / "model_architectures_with_dense_l3k24 (1)_annotated.pptx"
)

FONT = "Arial"
COLORS = {
    "ink": "25282D",
    "muted": "555555",
    "blue": "2F80ED",
    "blue_fill": "EAF2FD",
    "orange": "F07C42",
    "orange_fill": "FDF0E9",
    "green": "268A57",
    "green_fill": "E8F5EE",
    "red": "C74655",
    "red_fill": "FBECEF",
    "gray_fill": "F4F4F4",
    "line": "B8BEC6",
    "white": "FFFFFF",
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _name(shape, suffix: str) -> None:
    shape.name = f"GW_ANNOTATION_{suffix}"


def panel(slide, x, y, w, h, fill, line, suffix, radius=True, width=0.9):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    _name(shape, suffix)
    return shape


def textbox(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=9,
    color="25282D",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.02,
    suffix="TEXT",
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _name(shape, suffix)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 0.95
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def labeled_card(
    slide,
    x,
    y,
    w,
    h,
    label,
    body,
    *,
    accent,
    fill,
    suffix,
    label_size=7.8,
    body_size=7.1,
):
    panel(slide, x, y, w, h, fill, accent, f"{suffix}_PANEL", width=0.8)
    textbox(
        slide,
        label,
        x + 0.10,
        y + 0.05,
        w - 0.20,
        0.18,
        size=label_size,
        color=accent,
        bold=True,
        suffix=f"{suffix}_LABEL",
    )
    textbox(
        slide,
        body,
        x + 0.10,
        y + 0.23,
        w - 0.20,
        h - 0.27,
        size=body_size,
        color=COLORS["ink"],
        valign=MSO_ANCHOR.TOP,
        suffix=f"{suffix}_BODY",
    )


def compact_flag(slide, x, y, w, label, body, *, accent, fill, suffix):
    h = 0.58
    panel(slide, x, y, w, h, fill, accent, f"{suffix}_PANEL", width=0.8)
    textbox(
        slide,
        label,
        x + 0.08,
        y + 0.05,
        w - 0.16,
        0.16,
        size=7.3,
        color=accent,
        bold=True,
        suffix=f"{suffix}_LABEL",
    )
    textbox(
        slide,
        body,
        x + 0.08,
        y + 0.21,
        w - 0.16,
        0.31,
        size=6.8,
        color=COLORS["ink"],
        valign=MSO_ANCHOR.TOP,
        suffix=f"{suffix}_BODY",
    )


def shape_signature(shape):
    bbox = (shape.left, shape.top, shape.width, shape.height)
    text = getattr(shape, "text", "")
    table = None
    if getattr(shape, "has_table", False):
        table = tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
    return shape.shape_type, shape.name, bbox, text, table


def existing_signatures(prs):
    return [[shape_signature(shape) for shape in slide.shapes] for slide in prs.slides]


def annotate(prs: Presentation) -> None:
    if len(prs.slides) != 33:
        raise ValueError(f"Expected the 33-slide source deck, found {len(prs.slides)} slides")

    # Slide 9: a compact summary above the parameter ledger.  This distinguishes
    # graph/training choices from the component-wise parameter accounting below.
    slide = prs.slides[8]
    panel(slide, 0.55, 0.84, 12.20, 0.20, COLORS["blue_fill"], COLORS["blue"], "S09_STACK_PANEL")
    textbox(
        slide,
        "2.5° improvement stack  ·  row-aware kNN  +  L3 k24  +  Δ normalization  +  H160  +  S1×200  →  S2–S10",
        0.70,
        0.845,
        11.90,
        0.18,
        size=7.0,
        color=COLORS["blue"],
        bold=True,
        align=PP_ALIGN.CENTER,
        suffix="S09_STACK_TEXT",
    )

    # Slide 14: explain the observed 2.5-degree result improvement in the only
    # open band below the existing day-10 table and conclusion.
    slide = prs.slides[13]
    labeled_card(
        slide,
        0.50,
        6.42,
        3.99,
        0.64,
        "GRAPH · SHARED 2.5° STACK",
        "Row-aware kNN; coarse L3 is N=162, k=24 for broader mixing; per-channel Δ normalization.",
        accent=COLORS["blue"],
        fill=COLORS["blue_fill"],
        suffix="S14_GRAPH",
    )
    labeled_card(
        slide,
        4.67,
        6.42,
        3.99,
        0.64,
        "CAPACITY · 3.8M MODEL",
        "D 128→160; heads 4→5; 32/head fixed. Parameters 2.47M→3.84M (+55%).",
        accent=COLORS["orange"],
        fill=COLORS["orange_fill"],
        suffix="S14_CAPACITY",
    )
    labeled_card(
        slide,
        8.84,
        6.42,
        3.99,
        0.64,
        "TRAINING · 3.8M MODEL",
        "Strict 200-epoch S1 warm-start, then S2→S10 curriculum for 3 epochs per lead.",
        accent=COLORS["green"],
        fill=COLORS["green_fill"],
        suffix="S14_TRAINING",
    )
    textbox(
        slide,
        "Associated changes, not a controlled ablation: width and S1 pretraining duration changed together.",
        0.55,
        7.10,
        12.20,
        0.16,
        size=6.4,
        color=COLORS["muted"],
        align=PP_ALIGN.CENTER,
        suffix="S14_CAVEAT",
    )

    # Slide 30: connect the literature comparison to the actual ERA5-67 channel
    # inventory used by all experiments in this deck.
    slide = prs.slides[29]
    panel(slide, 0.75, 3.00, 11.80, 0.88, COLORS["gray_fill"], COLORS["line"], "S30_SCOPE_PANEL")
    textbox(
        slide,
        "OUR EXPERIMENT INPUT / OUTPUT CONTRACT",
        1.00,
        3.10,
        3.25,
        0.22,
        size=9.2,
        color=COLORS["ink"],
        bold=True,
        suffix="S30_SCOPE_LABEL",
    )
    textbox(
        slide,
        "ERA5-67 state: 60 upper-air fields (U/V/T/Q/Z × 12 levels) + t2m, msl, sp, tcwv, skt + TISR + orog.\nTwo input times → 134 channels; the head emits 67 channels. TISR and orog require special handling.",
        4.05,
        3.08,
        8.20,
        0.62,
        size=8.3,
        color=COLORS["ink"],
        suffix="S30_SCOPE_BODY",
    )

    # Slide 31: what the present dataset can and cannot support aloft.
    slide = prs.slides[30]
    compact_flag(
        slide,
        7.05,
        0.31,
        2.60,
        "USE IN OUR RUNS",
        "U / V / T / Q / Z at 1000…50 hPa (12 levels)",
        accent=COLORS["green"],
        fill=COLORS["green_fill"],
        suffix="S31_USE",
    )
    compact_flag(
        slide,
        9.82,
        0.31,
        2.98,
        "NOT IN ERA5-67",
        "Vertical velocity W — needs new data, stats, channels, and checkpoint",
        accent=COLORS["red"],
        fill=COLORS["red_fill"],
        suffix="S31_NO",
    )

    # Slide 32: available surface state versus fields that would require a
    # dataset/model-family extension.
    slide = prs.slides[31]
    compact_flag(
        slide,
        6.75,
        0.31,
        2.80,
        "USE IN OUR RUNS",
        "t2m · msl · sp · tcwv · skt",
        accent=COLORS["green"],
        fill=COLORS["green_fill"],
        suffix="S32_USE",
    )
    compact_flag(
        slide,
        9.72,
        0.31,
        3.08,
        "NOT IN ERA5-67",
        "10/100 m wind, TP, SST, soil, radiation, cloud, snow, runoff",
        accent=COLORS["red"],
        fill=COLORS["red_fill"],
        suffix="S32_NO",
    )

    # Slide 33: operationally safe use of the two non-prognostic channels and
    # explicit warning about features disabled in the shown checkpoints.
    slide = prs.slides[32]
    labeled_card(
        slide,
        0.50,
        5.34,
        3.35,
        1.26,
        "CAN USE NOW · INPUT/FORCING",
        "orog and TISR exist in every ERA5-67 file. Orography is static; TISR is deterministic at the forecast timestamp.",
        accent=COLORS["green"],
        fill=COLORS["green_fill"],
        suffix="S33_USE",
        label_size=8.2,
        body_size=8.0,
    )
    labeled_card(
        slide,
        4.00,
        5.34,
        4.05,
        1.26,
        "MUST NOT USE AS LEARNED FORECAST STATE",
        "Do not autoregress orog or TISR. Copy orog; inject the known future TISR; exclude both from loss. Never feed future ERA5 analysis fields.",
        accent=COLORS["red"],
        fill=COLORS["red_fill"],
        suffix="S33_NO_PREDICT",
        label_size=8.2,
        body_size=7.8,
    )
    labeled_card(
        slide,
        8.20,
        5.34,
        4.60,
        1.26,
        "NOT ENABLED IN THE SHOWN CHECKPOINTS",
        "Land–sea mask, explicit lat/lon/time inputs, and boundary forcing are absent. Adding them changes input width and requires a new checkpoint family.",
        accent=COLORS["blue"],
        fill=COLORS["blue_fill"],
        suffix="S33_NOT_ENABLED",
        label_size=8.2,
        body_size=7.8,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    prs = Presentation(args.input)
    original = existing_signatures(prs)
    annotate(prs)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)

    check = Presentation(args.output)
    if len(check.slides) != len(original):
        raise AssertionError("Slide count changed")
    for slide_index, (slide, signatures) in enumerate(zip(check.slides, original), start=1):
        if len(slide.shapes) < len(signatures):
            raise AssertionError(f"Slide {slide_index}: existing shapes were removed")
        current = [shape_signature(shape) for shape in list(slide.shapes)[: len(signatures)]]
        if current != signatures:
            raise AssertionError(f"Slide {slide_index}: existing content or geometry changed")

    added = sum(len(slide.shapes) - len(signatures) for slide, signatures in zip(check.slides, original))
    print(f"Wrote {args.output}")
    print(f"Slides preserved: {len(check.slides)}; existing shapes unchanged; annotations added: {added}")


if __name__ == "__main__":
    main()
