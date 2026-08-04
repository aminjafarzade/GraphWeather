#!/usr/bin/env python3
"""Add the 1.5-degree EMA qualitative maps to the architecture/results dashboard."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# Anchored at the repo root so the script works from any CWD and survives a
# checkout move (the decks now live in archive/).
REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE = REPO_ROOT / "archive" / "model_architectures_with_dense_l3k24.pptx"
OUTPUT = (
    REPO_ROOT
    / "archive"
    / "model_architectures_with_dense_l3k24_all_qualitative_predictions.pptx"
)
RUN_DIR = (
    REPO_ROOT
    / "runs"
    / "1p5_l4_h160_densel4k24_currS2toS10x3_initckpt_s1x150"
    / "visualizations_test_biasmaps_ema"
)

MAPS = [
    ("Z500", "500 hPa geopotential", RUN_DIR / "yearmean_z500_rollout_maps.png"),
    ("T2M", "2 m temperature", RUN_DIR / "yearmean_t2m_rollout_maps.png"),
    ("T850", "850 hPa temperature", RUN_DIR / "yearmean_t850_rollout_maps.png"),
    ("MSL", "mean sea-level pressure", RUN_DIR / "yearmean_msl_rollout_maps.png"),
]

VARIABLES = [
    ("Z500", "z500"),
    ("T2M", "t2m"),
    ("T850", "t850"),
    ("Q700", "q700"),
    ("U850", "u850"),
    ("MSL", "msl"),
]

H160_2P5_MAP_DIR = (
    REPO_ROOT
    / "runs"
    / "2p5_l3_h160_densel3k24_currS2toS10x3_initckpt_s1x200"
    / "visualizations_dashboard"
)
H128_2P5_MAP_DIR = (
    REPO_ROOT
    / "runs"
    / "2p5_l3_h128_densel3k24_currS2toS10x3_initckpt"
    / "visualizations_dashboard"
)


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    font: str = "Arial",
    color: str = "000000",
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_panel(slide, x, y, w, h, *, fill="F4F4F4", line="BBBBBB", width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def clear_slide(slide):
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)


def add_reference_title(slide, heading: str, subtitle: str):
    add_text(slide, heading, 0.50, 0.32, 12.30, 0.60, size=28, font="Cambria", bold=True)
    add_text(slide, subtitle, 0.50, 0.92, 12.30, 0.40, size=13, color="444444")


def add_prediction_panel(
    slide,
    path: Path,
    label: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    crop: tuple[float, float, float, float],
):
    add_panel(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill="FFFFFF", line="BBBBBB", width=0.8)
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom = crop
    add_text(slide, label, x + 0.08, y + 0.05, 0.82, 0.26, size=12, font="Cambria", bold=True)
    return picture


def rebuild_2p5_prediction_slide(slide, *, model_name: str, subtitle: str, map_dir: Path):
    clear_slide(slide)
    add_reference_title(slide, f"Qualitative predictions — {model_name}", subtitle)
    positions = [
        (0.55, 1.55), (4.73, 1.55), (8.91, 1.55),
        (0.55, 4.35), (4.73, 4.35), (8.91, 4.35),
    ]
    for (label, key), (x, y) in zip(VARIABLES, positions):
        path = map_dir / f"map_{key}_day10.png"
        add_prediction_panel(
            slide,
            path,
            label,
            x,
            y,
            3.87,
            2.26,
            crop=(1.0 / 3.0, 1.0 / 3.0, 0.0, 0.0),
        )


def rebuild_1p5_prediction_slide(slide):
    clear_slide(slide)
    add_reference_title(
        slide,
        "Qualitative predictions — GNN 4.6M EMA",
        "1.5° · year-mean day-10 prediction · fixed 10-day rollout · best EMA checkpoint · grid 121×240",
    )
    positions = [(0.67, 1.63), (7.10, 1.63), (0.67, 4.32), (7.10, 4.32)]
    for (short, _, path), (x, y) in zip(MAPS, positions):
        add_prediction_panel(
            slide,
            path,
            short,
            x,
            y,
            5.55,
            2.43,
            crop=(0.28, 0.35, 0.76, 0.03),
        )


def add_detail_slide(prs: Presentation, short: str, long_name: str, path: Path):
    # The Google-Slides-exported dashboard contains a single blank layout.
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_reference_title(
        slide,
        f"Qualitative rollout — {short}",
        f"1.5° GNN 4.6M EMA · {long_name} · GT, prediction, and bias at days 1, 3, 5, and 10",
    )

    add_panel(slide, 0.50, 1.55, 2.42, 2.20, fill="F4F4F4", line="BBBBBB")
    add_text(slide, "EXPERIMENT", 0.72, 1.76, 1.90, 0.24, size=10, font="Cambria", bold=True)
    add_text(
        slide,
        "1.5° L4 · hidden 160\ndense L4 k=24\nEMA best_ckpt\ngrid 121×240",
        0.72,
        2.08,
        1.92,
        1.30,
        size=11,
        color="444444",
        valign=MSO_ANCHOR.TOP,
    )

    add_panel(slide, 10.41, 1.55, 2.42, 2.20, fill="F4F4F4", line="BBBBBB")
    add_text(slide, "HOW TO READ", 10.63, 1.76, 1.90, 0.24, size=10, font="Cambria", bold=True)
    add_text(
        slide,
        "columns: GT | prediction | bias\nrows: days 1 | 3 | 5 | 10\nbias = prediction − GT\nRMSE shown per lead",
        10.63,
        2.08,
        1.92,
        1.30,
        size=10.5,
        color="444444",
        valign=MSO_ANCHOR.TOP,
    )

    # The image's natural aspect ratio is about 1.158, matching this frame.
    slide.shapes.add_picture(str(path), Inches(3.21), Inches(1.42), Inches(6.91), Inches(5.97))
    add_text(
        slide,
        "Year-mean diagnostics aggregate all test initializations; they reveal systematic spatial bias rather than one forecast case.",
        0.65,
        6.46,
        2.12,
        0.64,
        size=9.2,
        color="444444",
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "The same color scale is used for GT and prediction within each lead; bias has its own diverging scale.",
        10.56,
        6.46,
        2.12,
        0.64,
        size=9.2,
        color="444444",
        align=PP_ALIGN.CENTER,
    )
    return slide


def build():
    for _, _, path in MAPS:
        if not path.exists():
            raise FileNotFoundError(path)
    for map_dir in (H160_2P5_MAP_DIR, H128_2P5_MAP_DIR):
        for _, key in VARIABLES:
            path = map_dir / f"map_{key}_day10.png"
            if not path.exists():
                raise FileNotFoundError(path)
    prs = Presentation(SOURCE)
    if len(prs.slides) < 18:
        raise RuntimeError(f"Expected at least 18 slides, got {len(prs.slides)}")

    # Replace the three existing qualitative placeholder slides with actual
    # day-10 predictions, preserving their original experiment order.
    rebuild_2p5_prediction_slide(
        prs.slides[15],
        model_name="GNN 3.8M",
        subtitle="2.5° H160 · year-mean day-10 predictions · S1×200 warm start + S2→S10 curriculum · grid 72×144",
        map_dir=H160_2P5_MAP_DIR,
    )
    rebuild_2p5_prediction_slide(
        prs.slides[16],
        model_name="GNN 2.5M",
        subtitle="2.5° H128 · year-mean day-10 predictions · S2→S10 curriculum · grid 72×144",
        map_dir=H128_2P5_MAP_DIR,
    )
    rebuild_1p5_prediction_slide(prs.slides[17])
    insert_index = 18
    for short, long_name, path in MAPS:
        slide = add_detail_slide(prs, short, long_name, path)
        # Newly added slides start at the end; move them before the variable-comparison section.
        slide_ids = prs.slides._sldIdLst
        new_id = slide_ids[-1]
        slide_ids.remove(new_id)
        slide_ids.insert(insert_index, new_id)
        insert_index += 1

    prs.core_properties.title = "Model Architectures and Results with 1.5-degree Qualitative Maps"
    prs.core_properties.subject = "GraphWeather architectures, evaluation results, and EMA year-mean rollout maps"
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
